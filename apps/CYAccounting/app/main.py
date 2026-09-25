from __future__ import annotations

import faulthandler
import html
import os
import shutil
import sys
import traceback
from datetime import date, datetime, timedelta
from pathlib import Path


def _startup_root() -> Path:
    exe = Path(sys.executable).resolve()
    if exe.parent.name.lower() == "runtime":
        return exe.parent.parent
    return Path(__file__).resolve().parent.parent


_STARTUP_DATA_DIR = _startup_root() / "Data"
_STARTUP_DATA_DIR.mkdir(parents=True, exist_ok=True)
_STARTUP_LOG_PATH = _STARTUP_DATA_DIR / "startup_trace.log"
_STARTUP_LOG_MODE = "a" if os.environ.get("CY_STARTUP_TRACE_INITIALIZED") == "1" else "w"
_STARTUP_LOG_FILE = _STARTUP_LOG_PATH.open(_STARTUP_LOG_MODE, encoding="utf-8", buffering=1)


def startup_log(message: str) -> None:
    try:
        stamp = datetime.now().isoformat(timespec="milliseconds")
        _STARTUP_LOG_FILE.write(f"[{stamp}] {message}\n")
        _STARTUP_LOG_FILE.flush()
        os.fsync(_STARTUP_LOG_FILE.fileno())
    except Exception:
        pass


_LOG_DIR = _STARTUP_DATA_DIR / "Logs"
_LOG_DIR.mkdir(parents=True, exist_ok=True)

def _rotate_log_files(prefix: str, keep: int = 30) -> None:
    try:
        files = sorted(_LOG_DIR.glob(f"{prefix}_*.log"), key=lambda x: x.name, reverse=True)
        for old in files[keep:]:
            old.unlink(missing_ok=True)
    except Exception:
        pass

def app_log(message: str) -> None:
    try:
        path = _LOG_DIR / f"CYapp_{date.today().strftime('%Y%m%d')}.log"
        with path.open("a", encoding="utf-8") as handle:
            handle.write(f"[{datetime.now().isoformat(timespec='seconds')}] {message}\n")
        _rotate_log_files("CYapp")
    except Exception:
        pass

def error_log(message: str, exc_info=None) -> None:
    try:
        path = _LOG_DIR / f"CYerror_{date.today().strftime('%Y%m%d')}.log"
        with path.open("a", encoding="utf-8") as handle:
            handle.write("\n" + "=" * 72 + "\n")
            handle.write(f"[{datetime.now().isoformat(timespec='seconds')}] {message}\n")
            if exc_info:
                traceback.print_exception(*exc_info, file=handle)
        _rotate_log_files("CYerror")
    except Exception:
        pass

_SINGLE_INSTANCE_HANDLE = None

def acquire_single_instance() -> bool:
    global _SINGLE_INSTANCE_HANDLE
    if sys.platform != "win32":
        return True
    try:
        import ctypes
        kernel32 = ctypes.windll.kernel32
        user32 = ctypes.windll.user32
        handle = kernel32.CreateMutexW(None, False, "Local\\CYAccounting_SingleInstance")
        if not handle:
            return True
        if kernel32.GetLastError() == 183:  # ERROR_ALREADY_EXISTS
            hwnd = user32.FindWindowW(None, APP_NAME)
            if hwnd:
                user32.ShowWindow(hwnd, 9)
                user32.SetForegroundWindow(hwnd)
            ready = _STARTUP_DATA_DIR / "startup_ready.flag"
            ready.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
            startup_log("existing application instance detected; activated existing window")
            kernel32.CloseHandle(handle)
            return False
        _SINGLE_INSTANCE_HANDLE = handle
        return True
    except Exception as exc:
        startup_log(f"single instance check unavailable: {exc}")
        return True

def _early_exception_hook(exc_type, exc_value, exc_tb):
    startup_log(f"UNHANDLED EXCEPTION: {exc_type.__name__}: {exc_value}")
    try:
        traceback.print_exception(exc_type, exc_value, exc_tb, file=_STARTUP_LOG_FILE)
        _STARTUP_LOG_FILE.flush()
        os.fsync(_STARTUP_LOG_FILE.fileno())
    except Exception:
        pass


sys.excepthook = _early_exception_hook
startup_log(f"main.py entered | Python={sys.version.split()[0]} | executable={sys.executable}")
startup_log(f"cwd={Path.cwd()} | root={_startup_root()}")
try:
    faulthandler.enable(file=_STARTUP_LOG_FILE, all_threads=True)
    faulthandler.dump_traceback_later(15, repeat=False, file=_STARTUP_LOG_FILE)
    startup_log("faulthandler startup watchdog armed for 15 seconds")
except Exception as exc:
    startup_log(f"faulthandler unavailable: {exc}")

startup_log("importing PySide6")
from PySide6.QtCore import QDate, QEvent, QObject, QRect, QRegularExpression, QTimer, Qt
from PySide6.QtGui import QAction, QCloseEvent, QFont, QIcon, QIntValidator, QKeyEvent, QRegularExpressionValidator, QShortcut
from PySide6.QtWidgets import (
    QApplication,
    QAbstractItemView,
    QButtonGroup,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFrame,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QTabBar,
    QStackedWidget,
    QTableView,
    QToolButton,
    QVBoxLayout,
    QWidget,
)

from db import Database, DatabaseError
from background import background_tasks_running, start_background_task
from dialogs import (
    AccountManagerDialog,
    CarryForwardDialog,
    CategoryManagerDialog,
    EditTransactionDialog,
    OpeningBalanceDialog,
    SettingsDialog,
    no_tab_button,
)
from models import LedgerTableModel
from util import (
    APP_NAME,
    APP_VERSION,
    MAX_AMOUNT,
    BACKUP_PREFIX,
    DB_FILENAME,
    SaveResult,
    app_root,
    config_path,
    default_database_path,
    format_amount,
    format_date,
    is_month_locked,
    ConfigError,
    load_config_with_status,
    month_key_from_date,
    month_to_index,
    parse_date,
    save_config,
    safe_copy,
    shift_month,
    sync_version_marker,
    text_sort_key,
    try_set_zh_tw_locale,
    weighted_units,
)
from widgets import CategoryComboBox, DatePickerDialog, MonthSpinBox, SmartDateLineEdit, TransactionDelegate, WeightedLineEdit
from gdrive import GoogleDriveClient, GoogleDriveError, newest_local_backup, upload_backup_and_cleanup
from import_dialog import ImportTransactionsDialog

startup_log("all Python and PySide6 modules imported")


from theme import APP_STYLE


class ChineseStandardButtonFilter(QObject):
    """Force Qt standard dialog buttons to Traditional Chinese labels.

    The portable Qt runtime does not always include the zh-TW translation
    catalog, so convenience QMessageBox calls may otherwise show English
    labels such as OK, Cancel, Yes, and No.
    """

    DIALOG_LABELS = {
        QDialogButtonBox.StandardButton.Ok: "確定",
        QDialogButtonBox.StandardButton.Cancel: "取消",
        QDialogButtonBox.StandardButton.Save: "儲存",
        QDialogButtonBox.StandardButton.Close: "關閉",
        QDialogButtonBox.StandardButton.Apply: "套用",
        QDialogButtonBox.StandardButton.Yes: "是",
        QDialogButtonBox.StandardButton.No: "否",
        QDialogButtonBox.StandardButton.Retry: "重試",
        QDialogButtonBox.StandardButton.Ignore: "忽略",
        QDialogButtonBox.StandardButton.Abort: "中止",
        QDialogButtonBox.StandardButton.Discard: "放棄",
        QDialogButtonBox.StandardButton.Reset: "重設",
        QDialogButtonBox.StandardButton.RestoreDefaults: "恢復預設值",
    }
    MESSAGE_LABELS = {
        QMessageBox.StandardButton.Ok: "確定",
        QMessageBox.StandardButton.Cancel: "取消",
        QMessageBox.StandardButton.Yes: "是",
        QMessageBox.StandardButton.No: "否",
        QMessageBox.StandardButton.Save: "儲存",
        QMessageBox.StandardButton.Close: "關閉",
        QMessageBox.StandardButton.Apply: "套用",
        QMessageBox.StandardButton.Retry: "重試",
        QMessageBox.StandardButton.Ignore: "忽略",
        QMessageBox.StandardButton.Abort: "中止",
        QMessageBox.StandardButton.Discard: "放棄",
        QMessageBox.StandardButton.Reset: "重設",
        QMessageBox.StandardButton.RestoreDefaults: "恢復預設值",
    }

    @staticmethod
    def _localize_dialog_box(box: QDialogButtonBox) -> None:
        for standard, text in ChineseStandardButtonFilter.DIALOG_LABELS.items():
            button = box.button(standard)
            if button is not None:
                button.setText(text)

    @staticmethod
    def _localize_message_box(box: QMessageBox) -> None:
        for standard, text in ChineseStandardButtonFilter.MESSAGE_LABELS.items():
            button = box.button(standard)
            if button is not None:
                button.setText(text)

    @staticmethod
    def _clear_secondary_dialog_icon(dialog: QDialog) -> None:
        # CY Desktop visual rule: only the main application window carries the
        # canonical family icon. Keep native dialog chrome/behavior.
        dialog.setWindowIcon(QIcon())
        if sys.platform == "win32":
            try:
                import ctypes
                hwnd = int(dialog.winId())
                user32 = ctypes.windll.user32
                wm_seticon = 0x0080
                user32.SendMessageW(hwnd, wm_seticon, 0, 0)  # ICON_SMALL
                user32.SendMessageW(hwnd, wm_seticon, 1, 0)  # ICON_BIG
            except Exception:
                pass

    def eventFilter(self, watched, event):  # noqa: N802
        if event.type() in (QEvent.Type.Polish, QEvent.Type.Show):
            if isinstance(watched, QDialogButtonBox):
                self._localize_dialog_box(watched)
            elif isinstance(watched, QMessageBox):
                QTimer.singleShot(0, lambda box=watched: self._localize_message_box(box))
            elif isinstance(watched, QDialog) and not isinstance(watched, QFileDialog):
                QTimer.singleShot(0, lambda dlg=watched: self._clear_secondary_dialog_icon(dlg))
        return False


class InputTab(QWidget):
    data_changed = None

    def __init__(self, db: Database, config: dict, on_saved, parent=None):
        super().__init__(parent)
        self.setObjectName("inputPage")
        self.db = db
        self.config = config
        self.on_saved = on_saved
        self.confirm_lines: list[str] = []
        self._saving = False
        self.selected_account = ""
        self.account_buttons: dict[str, QPushButton] = {}
        self.quick_income_buttons: list[QPushButton] = []
        self.quick_expense_buttons: list[QPushButton] = []
        self.quick_income_summary_buttons: list[QPushButton] = []
        self.quick_expense_summary_buttons: list[QPushButton] = []
        self._build()
        self.refresh_lists(initial=True)
        self.date_edit.setText(format_date(date.today()))

    def _build(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(12, 6, 12, 8)
        outer.setSpacing(8)

        basic = QGroupBox("基本資訊")
        basic.setObjectName("inputSection")
        basic.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        basic_v = QVBoxLayout(basic)
        basic_v.setContentsMargins(10, 8, 10, 8)
        basic_v.setSpacing(8)
        basic_row = QHBoxLayout()
        basic_row.setContentsMargins(0, 0, 0, 0)
        basic_row.setSpacing(7)
        self.date_edit = SmartDateLineEdit()
        self.date_edit.setFixedWidth(145)
        self.date_edit.editingFinished.connect(self.validate_date)
        self.date_edit.textEdited.connect(lambda _text: self.date_error.hide() if hasattr(self, "date_error") else None)
        # Date controls share one visual height. Use Qt arrow primitives instead
        # of text glyphs so the day stepper stays crisp at Windows DPI scaling.
        date_control_height = max(34, self.date_edit.sizeHint().height())
        self.date_step_host = QWidget()
        self.date_step_host.setFixedSize(28, date_control_height)
        self.date_step_layout = QVBoxLayout(self.date_step_host)
        self.date_step_layout.setContentsMargins(0, 0, 0, 0)
        self.date_step_layout.setSpacing(1)
        upper_height = (date_control_height - 1) // 2
        lower_height = date_control_height - 1 - upper_height
        self.date_up_btn = QToolButton(self.date_step_host)
        self.date_up_btn.setObjectName("dateStepUp")
        self.date_up_btn.setArrowType(Qt.ArrowType.UpArrow)
        self.date_up_btn.setFixedSize(28, upper_height)
        self.date_up_btn.setToolTip("日期 +1 天")
        self.date_up_btn.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.date_up_btn.clicked.connect(lambda: self.shift_date(1))
        self.date_down_btn = QToolButton(self.date_step_host)
        self.date_down_btn.setObjectName("dateStepDown")
        self.date_down_btn.setArrowType(Qt.ArrowType.DownArrow)
        self.date_down_btn.setFixedSize(28, lower_height)
        self.date_down_btn.setToolTip("日期 -1 天")
        self.date_down_btn.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.date_down_btn.clicked.connect(lambda: self.shift_date(-1))
        self.date_step_layout.addWidget(self.date_up_btn)
        self.date_step_layout.addWidget(self.date_down_btn)

        self.calendar_btn = no_tab_button("")
        self.calendar_btn.setObjectName("calendarButton")
        self.calendar_btn.setIcon(QIcon(str(app_root() / "app" / "resources" / "calendar.png")))
        self.calendar_btn.setToolTip("選擇日期")
        self.calendar_btn.setFixedSize(42, date_control_height)
        self.calendar_btn.clicked.connect(self.pick_date)
        self.date_error = QLabel("日期錯誤")
        self.date_error.setStyleSheet("color:#B43737;font-weight:600;")
        self.date_error.setFixedWidth(76)
        self.date_error.hide()

        basic_row.addWidget(QLabel("日　　期："))
        basic_row.addWidget(self.date_edit)
        basic_row.addWidget(self.date_step_host)
        basic_row.addWidget(self.calendar_btn)
        basic_row.addWidget(self.date_error)
        basic_row.addSpacing(14)
        basic_row.addWidget(QLabel("帳　　戶："))
        self.account_button_host = QWidget()
        self.account_button_layout = QHBoxLayout(self.account_button_host)
        self.account_button_layout.setContentsMargins(0, 0, 0, 0)
        self.account_button_layout.setSpacing(6)
        basic_row.addWidget(self.account_button_host)
        basic_row.addStretch(1)
        basic_v.addLayout(basic_row)
        outer.addWidget(basic)

        income = QGroupBox("收入")
        income.setObjectName("incomeSection")
        income.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        income_v = QVBoxLayout(income)
        income_v.setContentsMargins(10, 8, 10, 8)
        income_v.setSpacing(8)
        income_quick = QHBoxLayout()
        income_quick.setSpacing(6)
        income_quick.addWidget(QLabel("常用科目："))
        self.income_quick_host = QWidget()
        self.income_quick_host.setFixedHeight(30)
        self.income_quick_layout = QHBoxLayout(self.income_quick_host)
        self.income_quick_layout.setContentsMargins(0, 0, 0, 0)
        self.income_quick_layout.setSpacing(5)
        income_quick.addWidget(self.income_quick_host)
        income_quick.addStretch(1)
        income_v.addLayout(income_quick)

        il = QHBoxLayout()
        il.setSpacing(7)
        self.income_category = CategoryComboBox()
        self.income_category.setFixedWidth(190)
        self.income_summary = WeightedLineEdit(40)
        self.income_summary.setPlaceholderText("可空白，最多20個中文字")
        self.income_amount = QLineEdit()
        self.income_amount.setFixedWidth(120)
        self.income_amount.setMaxLength(7)
        self.income_amount.setValidator(QRegularExpressionValidator(QRegularExpression(r"[0-9]{0,7}"), self.income_amount))
        self.income_amount.setPlaceholderText("最多7位")
        self.income_amount.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self.income_save = no_tab_button("存入收入")
        self.income_save.setFixedWidth(92)
        self.income_save.clicked.connect(lambda: self.save_entry("income"))
        il.addWidget(QLabel("收入科目："))
        il.addWidget(self.income_category)
        il.addSpacing(24)
        il.addWidget(QLabel("收入摘要："))
        il.addWidget(self.income_summary, 1)
        il.addWidget(QLabel("收入金額："))
        il.addWidget(self.income_amount)
        il.addWidget(self.income_save)
        income_v.addLayout(il)

        income_summary_band = QWidget()
        income_summary_band.setObjectName("summaryBand")
        income_summary_quick = QHBoxLayout(income_summary_band)
        income_summary_quick.setContentsMargins(0, 8, 0, 0)
        income_summary_quick.setSpacing(5)
        income_summary_quick.addWidget(QLabel("常用摘要："))
        self.income_quick_summary_host = QWidget()
        # Keep the common-summary row at a constant height whether it contains
        # buttons or the empty-state hint.  This prevents the whole input form
        # from shifting vertically when the selected category changes.
        self.income_quick_summary_host.setFixedHeight(28)
        self.income_quick_summary_layout = QHBoxLayout(self.income_quick_summary_host)
        self.income_quick_summary_layout.setContentsMargins(0, 3, 0, 3)
        self.income_quick_summary_layout.setSpacing(5)
        income_summary_quick.addWidget(self.income_quick_summary_host)
        income_summary_quick.addStretch(1)
        income_v.addWidget(income_summary_band)
        outer.addWidget(income)

        expense = QGroupBox("支出")
        expense.setObjectName("expenseSection")
        expense.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        expense_v = QVBoxLayout(expense)
        expense_v.setContentsMargins(10, 8, 10, 8)
        expense_v.setSpacing(8)
        expense_quick = QHBoxLayout()
        expense_quick.setSpacing(6)
        expense_quick.addWidget(QLabel("常用科目："))
        self.expense_quick_host = QWidget()
        self.expense_quick_host.setFixedHeight(30)
        self.expense_quick_layout = QHBoxLayout(self.expense_quick_host)
        self.expense_quick_layout.setContentsMargins(0, 0, 0, 0)
        self.expense_quick_layout.setSpacing(5)
        expense_quick.addWidget(self.expense_quick_host)
        expense_quick.addStretch(1)
        expense_v.addLayout(expense_quick)

        el = QHBoxLayout()
        el.setSpacing(7)
        self.expense_category = CategoryComboBox()
        self.expense_category.setFixedWidth(190)
        self.expense_summary = WeightedLineEdit(40)
        self.expense_summary.setPlaceholderText("可空白，最多20個中文字")
        self.expense_amount = QLineEdit()
        self.expense_amount.setFixedWidth(120)
        self.expense_amount.setMaxLength(7)
        self.expense_amount.setValidator(QRegularExpressionValidator(QRegularExpression(r"[0-9]{0,7}"), self.expense_amount))
        self.expense_amount.setPlaceholderText("最多7位")
        self.expense_amount.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self.expense_save = no_tab_button("存入支出")
        self.expense_save.setFixedWidth(92)
        self.expense_save.clicked.connect(lambda: self.save_entry("expense"))
        el.addWidget(QLabel("支出科目："))
        el.addWidget(self.expense_category)
        el.addSpacing(24)
        el.addWidget(QLabel("支出摘要："))
        el.addWidget(self.expense_summary, 1)
        el.addWidget(QLabel("支出金額："))
        el.addWidget(self.expense_amount)
        el.addWidget(self.expense_save)
        expense_v.addLayout(el)

        expense_summary_band = QWidget()
        expense_summary_band.setObjectName("summaryBand")
        expense_summary_quick = QHBoxLayout(expense_summary_band)
        expense_summary_quick.setContentsMargins(0, 8, 0, 0)
        expense_summary_quick.setSpacing(5)
        expense_summary_quick.addWidget(QLabel("常用摘要："))
        self.expense_quick_summary_host = QWidget()
        self.expense_quick_summary_host.setFixedHeight(28)
        self.expense_quick_summary_layout = QHBoxLayout(self.expense_quick_summary_host)
        self.expense_quick_summary_layout.setContentsMargins(0, 3, 0, 3)
        self.expense_quick_summary_layout.setSpacing(5)
        expense_summary_quick.addWidget(self.expense_quick_summary_host)
        expense_summary_quick.addStretch(1)
        expense_v.addWidget(expense_summary_band)
        outer.addWidget(expense)

        confirm = QGroupBox("輸入確認")
        confirm.setObjectName("confirmationCard")
        confirm.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        confirm.setFixedHeight(244)
        cv = QVBoxLayout(confirm)
        cv.setContentsMargins(10, 7, 10, 7)
        cv.setSpacing(2)
        self.confirm_labels = []
        for _ in range(10):
            lab = QLabel("")
            lab.setMinimumHeight(18)
            lab.setTextFormat(Qt.TextFormat.RichText)
            self.confirm_labels.append(lab)
            cv.addWidget(lab)
        outer.addWidget(confirm)
        outer.addStretch(1)

        self._static_main_fields = [
            self.date_edit, self.income_category, self.income_summary, self.income_amount,
            self.expense_category, self.expense_summary, self.expense_amount,
        ]
        for w in self._static_main_fields:
            w.installEventFilter(self)
        self.income_category.currentIndexChanged.connect(lambda *_: self._build_quick_summary_buttons("income"))
        self.expense_category.currentIndexChanged.connect(lambda *_: self._build_quick_summary_buttons("expense"))
        QTimer.singleShot(0, self._sync_entry_action_heights)

    def _sync_entry_action_heights(self):
        for field, button in (
            (self.income_amount, self.income_save),
            (self.expense_amount, self.expense_save),
        ):
            target_height = max(field.height(), field.sizeHint().height())
            if target_height > 0:
                button.setFixedHeight(target_height)

    def _clear_layout(self, layout: QHBoxLayout):
        while layout.count():
            item = layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()

    def _build_account_buttons(self, initial: bool = False):
        previous = self.selected_account
        accounts = [a["name"] for a in self.db.accounts()]
        target = self.db.default_account() if initial or previous not in accounts else previous
        self._clear_layout(self.account_button_layout)
        self.account_buttons = {}
        group = QButtonGroup(self.account_button_host)
        group.setExclusive(True)
        self._account_button_group = group
        for name in accounts:
            button = QPushButton(name)
            button.setObjectName("accountChoiceButton")
            button.setCheckable(True)
            button.setMinimumWidth(78)
            button.setMaximumWidth(132)
            button.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            button.clicked.connect(lambda checked=False, n=name: self.select_account(n, focus=False))
            button.installEventFilter(self)
            group.addButton(button)
            self.account_button_layout.addWidget(button)
            self.account_buttons[name] = button
        self.selected_account = target if target in self.account_buttons else (accounts[0] if accounts else "")
        if self.selected_account:
            self._set_account_visual(self.selected_account)
        self._rewire_tab_order()

    def _set_account_visual(self, name: str):
        for account_name, button in self.account_buttons.items():
            selected = account_name == name
            button.setChecked(selected)
            button.setToolTip("目前使用中的帳戶" if selected else f"切換至帳戶「{account_name}」")
            button.setFocusPolicy(Qt.FocusPolicy.StrongFocus if selected else Qt.FocusPolicy.NoFocus)
        self.selected_account = name

    def select_account(self, name: str, focus: bool = False):
        if name not in self.account_buttons:
            return
        self._set_account_visual(name)
        self._rewire_tab_order()
        # Common summaries are scoped by account + kind + category.
        self._build_quick_summary_buttons("income")
        self._build_quick_summary_buttons("expense")
        if focus:
            self.account_buttons[name].setFocus()

    def _build_quick_buttons(self, kind: str):
        layout = self.income_quick_layout if kind == "income" else self.expense_quick_layout
        self._clear_layout(layout)
        buttons = []
        for name in self.db.favorite_category_names(kind):
            button = no_tab_button(name)
            button.setObjectName("quickCategoryButton")
            button.setMinimumHeight(26)
            button.setToolTip(f"切換{('收入' if kind == 'income' else '支出')}科目為「{name}」")
            if kind == "income":
                button.clicked.connect(lambda checked=False, n=name: self.income_category.set_current_value(n))
            else:
                button.clicked.connect(lambda checked=False, n=name: self.expense_category.set_current_value(n))
            layout.addWidget(button)
            buttons.append(button)
        if not buttons:
            hint = QLabel("尚未設定")
            hint.setStyleSheet("color:#8a94a0;")
            layout.addWidget(hint)
        if kind == "income":
            self.quick_income_buttons = buttons
        else:
            self.quick_expense_buttons = buttons

    @staticmethod
    def _summary_button_text(summary: str) -> str:
        text = summary.strip()
        return text if len(text) <= 8 else text[:8] + "…"

    def _build_quick_summary_buttons(self, kind: str):
        if not hasattr(self, "income_quick_summary_layout"):
            return
        layout = self.income_quick_summary_layout if kind == "income" else self.expense_quick_summary_layout
        self._clear_layout(layout)
        category = self.income_category.current_value() if kind == "income" else self.expense_category.current_value()
        recent_n = int(self.config.get("common_summary_recent_n", 100) or 100)
        min_count = int(self.config.get("common_summary_min_count", 3) or 3)
        basis = self.config.get("common_summary_basis", "tx_date") or "tx_date"
        summaries = (
            self.db.frequent_summaries(
                kind, self.selected_account, category, recent_n, min_count, 10, basis
            )
            if category and self.selected_account else []
        )
        buttons: list[QPushButton] = []
        target_edit = self.income_summary if kind == "income" else self.expense_summary
        amount_edit = self.income_amount if kind == "income" else self.expense_amount
        for summary in summaries:
            button = no_tab_button(self._summary_button_text(summary))
            button.setObjectName("quickSummaryButton")
            button.setToolTip(summary)
            button.clicked.connect(
                lambda checked=False, s=summary, edit=target_edit, amount=amount_edit: (
                    edit.setText(s), amount.setFocus(), amount.selectAll()
                )
            )
            layout.addWidget(button)
            buttons.append(button)
        if not buttons:
            hint = QLabel("尚無符合條件")
            hint.setFixedHeight(22)
            hint.setAlignment(Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft)
            hint.setStyleSheet("color:#98A2B3;")
            layout.addWidget(hint)
        if kind == "income":
            self.quick_income_summary_buttons = buttons
        else:
            self.quick_expense_summary_buttons = buttons

    def _rewire_tab_order(self):
        account_button = self.account_buttons.get(self.selected_account)
        if account_button is not None:
            QWidget.setTabOrder(self.date_edit, account_button)
            QWidget.setTabOrder(account_button, self.income_category)
        else:
            QWidget.setTabOrder(self.date_edit, self.income_category)
        QWidget.setTabOrder(self.income_category, self.income_summary)
        QWidget.setTabOrder(self.income_summary, self.income_amount)
        QWidget.setTabOrder(self.income_amount, self.expense_category)
        QWidget.setTabOrder(self.expense_category, self.expense_summary)
        QWidget.setTabOrder(self.expense_summary, self.expense_amount)
        QWidget.setTabOrder(self.expense_amount, self.date_edit)

    def refresh_lists(self, initial: bool = False):
        inc = self.income_category.current_value() if hasattr(self, "income_category") else ""
        exp = self.expense_category.current_value() if hasattr(self, "expense_category") else ""
        self._build_account_buttons(initial=initial)
        self.income_category.set_tree(self.db.category_tree("income"), inc or None)
        self.expense_category.set_tree(self.db.category_tree("expense"), exp or None)
        self._build_quick_buttons("income")
        self._build_quick_buttons("expense")
        self._build_quick_summary_buttons("income")
        self._build_quick_summary_buttons("expense")

    def validate_date(self) -> bool:
        self.date_edit.normalize_format()
        valid = parse_date(self.date_edit.text()) is not None
        self.date_error.setVisible(not valid)
        if not valid:
            self.date_edit.setFocus(); self.date_edit.selectAll()
        return valid

    def pick_date(self):
        initial = parse_date(self.date_edit.text()) or date.today()
        dlg = DatePickerDialog(initial, self)
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.selected_date:
            self.date_edit.setText(format_date(dlg.selected_date))
            self.date_error.hide()

    def shift_date(self, delta: int):
        d = parse_date(self.date_edit.text())
        if d is None:
            self.validate_date(); return
        self.date_edit.setText(format_date(d + timedelta(days=delta)))
        self.date_error.hide()

    @staticmethod
    def _plain_or_keypad(mods) -> bool:
        # The numeric keypad Enter carries KeypadModifier.  Treat it exactly
        # like the main Enter while still rejecting Ctrl/Shift/Alt/Meta combos.
        blocked = (
            Qt.KeyboardModifier.ShiftModifier
            | Qt.KeyboardModifier.ControlModifier
            | Qt.KeyboardModifier.AltModifier
            | Qt.KeyboardModifier.MetaModifier
        )
        return not bool(mods & blocked)

    def eventFilter(self, watched, event):  # noqa: N802
        if event.type() == QEvent.Type.KeyPress:
            key = event.key()
            mods = event.modifiers()
            if mods & Qt.KeyboardModifier.ControlModifier:
                if key == Qt.Key.Key_Up:
                    self.shift_date(1); return True
                if key == Qt.Key.Key_Down:
                    self.shift_date(-1); return True
            if self._plain_or_keypad(mods):
                is_enter = key in (Qt.Key.Key_Return, Qt.Key.Key_Enter)
                if watched is self.date_edit and is_enter:
                    if self.validate_date():
                        account_button = self.account_buttons.get(self.selected_account)
                        if account_button:
                            account_button.setFocus()
                        else:
                            self.income_category.setFocus()
                    return True
                if watched in self.account_buttons.values() and is_enter:
                    self.income_category.setFocus(); return True
                if watched is self.income_category and is_enter:
                    if self.income_category.view().isVisible():
                        return False
                    self.income_summary.setFocus(); return True
                if watched is self.expense_category and is_enter:
                    if self.expense_category.view().isVisible():
                        return False
                    self.expense_summary.setFocus(); return True
                if watched is self.income_summary and key == Qt.Key.Key_Down:
                    self.expense_summary.setFocus(); return True
                if watched is self.expense_summary and key == Qt.Key.Key_Up:
                    self.income_summary.setFocus(); return True
                if watched is self.income_amount and key == Qt.Key.Key_Down:
                    self.expense_amount.setFocus(); return True
                if watched is self.expense_amount and key == Qt.Key.Key_Up:
                    self.income_amount.setFocus(); return True
                if watched is self.income_summary and is_enter:
                    self.income_amount.setFocus(); self.income_amount.selectAll(); return True
                if watched is self.expense_summary and is_enter:
                    self.expense_amount.setFocus(); self.expense_amount.selectAll(); return True
                if watched is self.income_amount and is_enter:
                    self.save_entry("income"); return True
                if watched is self.expense_amount and is_enter:
                    self.save_entry("expense"); return True
        return super().eventFilter(watched, event)

    def _add_confirmation(self, html_line: str):
        self.confirm_lines.insert(0, html_line)
        self.confirm_lines = self.confirm_lines[:10]
        for i, label in enumerate(self.confirm_labels):
            label.setText(self.confirm_lines[i] if i < len(self.confirm_lines) else "")

    def show_failure(self, reason: str):
        self._add_confirmation(
            f'<span style="color:#B43737;font-weight:bold;">存檔失敗</span>　'
            f'<span style="color:#B43737;font-size:10pt;">{html.escape(reason)}</span>'
        )

    def _auto_carry_if_needed(self, d: date) -> dict[str, int] | None:
        month = month_key_from_date(d)
        if d.day != 1 or self.db.month_has_transactions(month) or self.db.month_has_any_opening(month):
            return None
        values = self.db.previous_month_carry_values(month)
        if values is None:
            return None
        dlg = CarryForwardDialog(month, shift_month(month, -1), values, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            return values
        return None

    def save_entry(self, kind: str):
        if self._saving:
            return
        self._saving = True
        try:
            d = parse_date(self.date_edit.text())
            if d is None:
                self.date_error.show(); self.date_edit.setFocus(); self.date_edit.selectAll()
                self.show_failure("日期格式或日期內容不正確")
                return
            self.date_error.hide()
            account = self.selected_account.strip()
            category = self.income_category.current_value() if kind == "income" else self.expense_category.current_value()
            summary_edit = self.income_summary if kind == "income" else self.expense_summary
            amount_edit = self.income_amount if kind == "income" else self.expense_amount
            if not account:
                self.show_failure("尚未選擇帳戶"); return
            if not category:
                self.show_failure("尚未選擇科目"); return
            if weighted_units(summary_edit.text()) > 40:
                summary_edit.setFocus(); self.show_failure("摘要不可超過20個中文字或40個英數字元"); return
            try:
                amount = int(amount_edit.text())
            except Exception:
                amount = 0
            if amount < 1:
                amount_edit.setFocus(); amount_edit.selectAll(); self.show_failure("金額必須大於0"); return
            if amount > MAX_AMOUNT:
                amount_edit.setFocus(); amount_edit.selectAll(); self.show_failure("金額最多7位數"); return
            if is_month_locked(month_key_from_date(d), self.db.locked_through()):
                self.show_failure(f"{month_key_from_date(d)} 已鎖定，無法新增資料"); return
            carry_values = self._auto_carry_if_needed(d)
            result = self.db.save_transaction(
                format_date(d), account, kind, category, summary_edit.text(), amount,
                opening_values=carry_values,
            )
            if not result.ok:
                self.show_failure(result.message)
                return
            type_text = "收入" if kind == "income" else "支出"
            if kind == "income":
                type_tag = (
                    '<span style="background-color:#EDF5F2;color:#2E6F5E;">'
                    '&nbsp;收入&nbsp;</span>'
                )
            else:
                type_tag = (
                    '<span style="background-color:#F6EEEE;color:#8A5B5B;">'
                    '&nbsp;支出&nbsp;</span>'
                )
            summary_text = summary_edit.text().strip()
            summary_display = (
                html.escape(summary_text)
                if summary_text
                else '<span style="color:#98A2B3;">(空白)</span>'
            )
            # Confirmation layout uses full-width spaces as visual separators.
            # The income/expense tag replaces the old "收入-科目 / 支出-科目" text.
            record = (
                f"{format_date(d)}　[{html.escape(account)}]　{type_tag}　{html.escape(category)}　-　"
                f"{summary_display}　${amount:,}　"
                f'<span style="color:#21825C;font-weight:bold;">&lt;存檔成功&gt;</span>'
            )
            self._add_confirmation(record)
            summary_edit.clear(); amount_edit.clear(); summary_edit.setFocus()
            self._build_quick_summary_buttons(kind)
            self.on_saved(month_key_from_date(d))
        finally:
            self._saving = False

    def open_account_manager(self):
        old = self.selected_account
        AccountManagerDialog(self.db, self).exec()
        self.refresh_lists()
        if old in self.account_buttons:
            self.select_account(old)

    def open_category_manager(self):
        inc = self.income_category.current_value(); exp = self.expense_category.current_value()
        CategoryManagerDialog(self.db, self).exec()
        self.income_category.set_tree(self.db.category_tree("income"), inc)
        self.expense_category.set_tree(self.db.category_tree("expense"), exp)
        self._build_quick_buttons("income")
        self._build_quick_buttons("expense")
        self._build_quick_summary_buttons("income")
        self._build_quick_summary_buttons("expense")

    def has_unsaved_content(self) -> bool:
        return any((self.income_summary.text(), self.income_amount.text(), self.expense_summary.text(), self.expense_amount.text()))


class LedgerTab(QWidget):
    def __init__(self, db: Database, initial_month: str, config: dict | None = None, parent=None):
        super().__init__(parent)
        self.db = db
        self.config = config if config is not None else {}
        self.displayed_month = initial_month
        self.current_data: dict = {}
        self.search_query = ""
        self.account_sort_mode = False
        self._pending_reselect: int | None = None
        self._build()
        self.month_spin.set_month(initial_month)
        self.load_month(initial_month)

    def _build(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(12, 6, 12, 8)
        outer.setSpacing(5)
        top = QHBoxLayout()
        top.setSpacing(6)
        top.addWidget(QLabel("選擇月份："))
        self.month_spin = MonthSpinBox()
        top.addWidget(self.month_spin)
        self.query_btn = no_tab_button("查詢"); self.query_btn.setObjectName("primaryButton"); self.query_btn.clicked.connect(self.query_selected)
        top.addWidget(self.query_btn)
        top.addSpacing(8)
        self.search_edit = QLineEdit()
        self.search_edit.setFixedWidth(220)
        self.search_edit.setPlaceholderText("搜尋收入／支出摘要")
        self.search_edit.returnPressed.connect(self.search_current_month)
        top.addWidget(self.search_edit)
        self.search_btn = no_tab_button("搜尋")
        self.search_btn.clicked.connect(self.search_current_month)
        top.addWidget(self.search_btn)
        top.addStretch()
        self.opening_btn = no_tab_button("設定期初餘額")
        self.opening_btn.clicked.connect(self.set_opening)
        top.addWidget(self.opening_btn)
        outer.addLayout(top)

        header_row = QHBoxLayout()
        header_row.setContentsMargins(0, 0, 0, 0)
        header_row.setSpacing(6)
        self.table_title = QLabel()
        f = self.table_title.font(); f.setBold(True); f.setPointSize(f.pointSize()+1); self.table_title.setFont(f)
        header_row.addWidget(self.table_title)
        self.search_status = QLabel("")
        self.search_status.setStyleSheet("color:#667085;")
        header_row.addWidget(self.search_status)
        header_row.addStretch()
        self.opening_warning = QLabel("此月份尚有帳戶未設定期初餘額")
        self.opening_warning.setStyleSheet("color:#B43737;font-weight:bold;")
        header_row.addWidget(self.opening_warning)
        outer.addLayout(header_row)

        self.model = LedgerTableModel(self.inline_edit)
        self.model.edit_failed.connect(lambda msg: QMessageBox.warning(self, APP_NAME, msg))
        self.model.edit_succeeded.connect(self.reload_reselect)
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QTableView.SelectionMode.ExtendedSelection)
        self.table.setEditTriggers(QTableView.EditTrigger.DoubleClicked)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setMinimumSectionSize(22)
        self.table.verticalHeader().setDefaultSectionSize(23)
        table_font = self.table.font()
        table_font.setPointSize(10)
        self.table.setFont(table_font)
        header_font = self.table.horizontalHeader().font()
        header_font.setPointSize(10)
        header_font.setBold(True)
        self.table.horizontalHeader().setFont(header_font)
        self.table.horizontalHeader().setFixedHeight(24)
        header = self.table.horizontalHeader()
        header.setMinimumSectionSize(55)
        header.setStretchLastSection(False)
        header.setSectionsMovable(False)
        header.setSectionsClickable(True)
        header.sectionClicked.connect(self._header_clicked)

        # V1.0.21: fixed, non-draggable column layout.  The fixed-width
        # columns total 736 px at the normal DPI; the two summary columns
        # share all remaining viewport width.  This keeps income/expense
        # amount columns identical, gives the balance column a little more
        # room, and prevents old saved widths from distorting the layout.
        self._fixed_column_widths = {
            0: 92,   # 日期
            1: 100,  # 帳戶
            2: 60,   # 收支
            3: 105,  # 收入科目
            5: 92,   # 收入金額
            6: 105,  # 支出科目
            8: 92,   # 支出金額（與收入金額相同）
            9: 90,   # 餘額（V1.0.20 的 74 px 稍微加寬）
        }
        for col in range(self.model.columnCount()):
            if col in (4, 7):
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.Stretch)
            else:
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.Fixed)
                self.table.setColumnWidth(col, self._fixed_column_widths[col])
        self.table.setItemDelegate(TransactionDelegate(
            lambda: [a["name"] for a in self.db.accounts()],
            lambda kind: self.db.category_names(kind), self.table
        ))
        self.table.selectionModel().selectionChanged.connect(self.update_action_state)
        outer.addWidget(self.table, 1)

        bottom = QHBoxLayout()
        bottom.setSpacing(8)

        stats_box = QGroupBox("月份統計")
        stats_box.setObjectName("statsCard")
        stats_box.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        sg = QGridLayout(stats_box)
        sg.setContentsMargins(10, 6, 10, 6)
        sg.setHorizontalSpacing(24)
        sg.setVerticalSpacing(2)
        self.opening_label = QLabel(); self.income_label = QLabel(); self.expense_label = QLabel(); self.net_label = QLabel(); self.ending_label = QLabel()
        labels = [self.opening_label, self.income_label, self.expense_label, self.net_label, self.ending_label]
        for lab in labels:
            font = lab.font(); font.setBold(True); lab.setFont(font)
        self.income_label.setStyleSheet("color:#2E6F5E;font-weight:bold;")
        self.expense_label.setStyleSheet("color:#8A5B5B;font-weight:bold;")
        self.ending_label.setStyleSheet("color:#1F2937;font-weight:bold;")
        sg.addWidget(self.opening_label, 0, 0, 1, 2)
        sg.addWidget(self.income_label, 1, 0); sg.addWidget(self.expense_label, 1, 1)
        sg.addWidget(self.net_label, 2, 0); sg.addWidget(self.ending_label, 2, 1)
        self.lock_btn = no_tab_button("確認並鎖定本月")
        self.lock_btn.clicked.connect(self.confirm_lock)
        self.lock_btn.setMinimumWidth(150)
        sg.addWidget(self.lock_btn, 0, 2, 3, 1, Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        sg.setColumnStretch(0, 1)
        sg.setColumnStretch(1, 1)
        bottom.addWidget(stats_box, 1)

        action_panel = QWidget()
        action_panel.setFixedWidth(158)
        action_col = QVBoxLayout(action_panel)
        action_col.setContentsMargins(0, 0, 0, 0)
        action_col.setSpacing(5)
        self.edit_btn = no_tab_button("編輯選取單筆"); self.edit_btn.clicked.connect(self.edit_selected)
        self.delete_btn = no_tab_button("刪除選取項目"); self.delete_btn.setObjectName("dangerButton"); self.delete_btn.clicked.connect(self.delete_selected)
        self.export_btn = no_tab_button("單月匯出 Excel"); self.export_btn.clicked.connect(self.export_excel)
        for button in (self.edit_btn, self.delete_btn, self.export_btn):
            button.setMinimumWidth(158)
            action_col.addWidget(button)
        action_col.addStretch(1)
        bottom.addWidget(action_panel, 0, Qt.AlignmentFlag.AlignBottom)
        outer.addLayout(bottom)

    def resizeEvent(self, event):  # noqa: N802
        super().resizeEvent(event)
        if hasattr(self, "table"):
            rows = max(6, int(self.table.viewport().height() / max(1, self.table.verticalHeader().defaultSectionSize())))
            if rows != self.model.placeholder_rows:
                self.model.set_rows(self.model.rows, rows)

    def query_selected(self):
        try:
            month = self.month_spin.month()
        except ValueError:
            return
        self.search_edit.clear()
        self.search_query = ""
        self._set_account_sort_mode(False)
        self.load_month(month)

    def search_current_month(self):
        self.search_query = self.search_edit.text().strip()
        self._apply_summary_filter()

    def _filtered_rows(self) -> list[dict]:
        rows = list(self.current_data.get("rows", []))
        query = self.search_query.casefold()
        if query:
            rows = [row for row in rows if query in (row.get("summary") or "").casefold()]
        if self.account_sort_mode:
            rows.sort(key=lambda r: (
                text_sort_key(r.get("account_name", "")),
                r.get("tx_date", ""),
                0 if r.get("kind") == "income" else 1,
                r.get("created_at", ""),
                int(r.get("id", 0)),
            ))
        return rows

    def _set_account_sort_mode(self, enabled: bool) -> None:
        enabled = bool(enabled)
        if self.account_sort_mode == enabled:
            return
        self.account_sort_mode = enabled
        self.model.set_account_view(enabled)
        self._apply_summary_filter()

    def _header_clicked(self, section: int) -> None:
        if section == 1:
            self._set_account_sort_mode(not self.account_sort_mode)

    def _apply_summary_filter(self, reselect_id: int | None = None):
        rows = self._filtered_rows()
        placeholder_rows = max(6, int(self.table.viewport().height() / max(1, self.table.verticalHeader().defaultSectionSize())))
        self.model.set_rows(rows, placeholder_rows)
        if self.search_query:
            if rows:
                self.search_status.setText(f"搜尋結果：{len(rows)} 筆")
                self.search_status.setStyleSheet("color:#667085;")
            else:
                self.search_status.setText("查無符合摘要資料")
                self.search_status.setStyleSheet("color:#B43737;font-weight:bold;")
        else:
            self.search_status.clear()
            self.search_status.setStyleSheet("color:#667085;")
        self.update_action_state()
        if reselect_id is not None:
            for row_index, data in enumerate(self.model.rows):
                if int(data["id"]) == int(reselect_id):
                    self.table.selectRow(row_index)
                    self.table.scrollTo(self.model.index(row_index, 0))
                    break

    def load_month(self, month: str, reselect_id: int | None = None):
        self.displayed_month = month
        self.current_data = self.db.month_data(month)
        locked = is_month_locked(month, self.db.locked_through())
        self.model.set_locked(locked)
        self._apply_summary_filter(reselect_id)
        self.table_title.setText(f"目前顯示｜{month}" + ("　🔒 已鎖定" if locked else ""))
        self.opening_warning.setVisible(bool(self.current_data["missing_openings"]))
        self.opening_label.setText(f"期初餘額：{format_amount(self.current_data['opening_total'])}")
        self.income_label.setText(f"收入小計：{format_amount(self.current_data['income_total'])}")
        self.expense_label.setText(f"支出小計：{format_amount(self.current_data['expense_total'])}")
        net = self.current_data["net"]
        if net >= 0:
            self.net_label.setText(f"淨利：{format_amount(net)}")
            self.net_label.setStyleSheet("color:#21825C;font-weight:bold;")
        else:
            self.net_label.setText(f"淨損：{format_amount(abs(net))}")
            self.net_label.setStyleSheet("color:#8A5B5B;font-weight:bold;")
        self.ending_label.setText(f"期末餘額：{format_amount(self.current_data['ending_total'])}")
        opening_tip = self._balances_tooltip(self.current_data["openings"], self.current_data["opening_total"], self.current_data["missing_openings"])
        ending_tip = self._balances_tooltip(self.current_data["ending_by_account"], self.current_data["ending_total"], [])
        self.opening_label.setToolTip(opening_tip); self.ending_label.setToolTip(ending_tip)
        self.opening_btn.setEnabled(not locked)
        self._update_lock_button()
        self.update_action_state()

    def update_action_state(self, *_):
        locked = is_month_locked(self.displayed_month, self.db.locked_through())
        count = len(self.selected_ids()) if hasattr(self, "table") else 0
        self.edit_btn.setEnabled((not locked) and count == 1)
        self.delete_btn.setEnabled((not locked) and count >= 1)

    def column_widths(self) -> list[int]:
        return [self.table.columnWidth(i) for i in range(self.model.columnCount())]

    def restore_column_widths(self, widths) -> None:
        # V1.0.21: column widths are intentionally fixed by the application.
        # Ignore legacy saved widths so users cannot end up with an oversized
        # amount column or a squeezed balance column after an upgrade.
        header = self.table.horizontalHeader()
        for col in range(self.model.columnCount()):
            if col in (4, 7):
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.Stretch)
            else:
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.Fixed)
                self.table.setColumnWidth(col, self._fixed_column_widths[col])

    def scroll_to_entry_position(self, position: str) -> None:
        """Scroll to the oldest/latest real transaction row without selecting it."""
        row_count = len(self.model.rows)
        if row_count <= 0:
            self.table.scrollToTop()
            return
        if position == "oldest":
            index = self.model.index(0, 0)
            self.table.scrollTo(index, QAbstractItemView.ScrollHint.PositionAtTop)
        else:
            index = self.model.index(row_count - 1, 0)
            self.table.scrollTo(index, QAbstractItemView.ScrollHint.PositionAtBottom)

    def _balances_tooltip(self, values: dict[str, int], total: int, missing: list[str]) -> str:
        names = [name for name, _ in self.db.relevant_accounts_for_month(self.displayed_month)]
        for n in values:
            if n not in names:
                names.append(n)
        lines = []
        for name in names:
            suffix = "（未設定）" if name in missing else ""
            lines.append(f"{name}：{format_amount(values.get(name, 0))}{suffix}")
        lines.append("────────────")
        lines.append(f"總餘額：{format_amount(total)}")
        return "\n".join(lines)

    def _update_lock_button(self):
        locked = self.db.locked_through()
        if locked:
            candidate = shift_month(locked, 1)
        else:
            candidate = self.db.earliest_data_month()
        self.lock_btn.setVisible(bool(candidate and candidate == self.displayed_month and not is_month_locked(self.displayed_month, locked)))

    def inline_edit(self, tx_id: int, field: str, value) -> tuple[bool, str]:
        return self.db.update_transaction(tx_id, field, value)

    def reload_reselect(self, tx_id: int):
        tx = self.db.get_transaction(tx_id)
        if tx and tx["tx_date"][:7] == self.displayed_month:
            self.load_month(self.displayed_month, tx_id)
        else:
            self.load_month(self.displayed_month)

    def selected_ids(self) -> list[int]:
        ids = []
        for idx in self.table.selectionModel().selectedRows():
            row = self.model.row_data(idx.row())
            if row:
                ids.append(int(row["id"]))
        return list(dict.fromkeys(ids))

    def set_opening(self):
        if is_month_locked(self.displayed_month, self.db.locked_through()):
            return
        if OpeningBalanceDialog(self.db, self.displayed_month, self).exec() == QDialog.DialogCode.Accepted:
            self.load_month(self.displayed_month)

    def edit_selected(self):
        ids = self.selected_ids()
        if len(ids) != 1:
            QMessageBox.information(self, APP_NAME, "請先選取一筆資料。")
            return
        tx = self.db.get_transaction(ids[0])
        if tx and EditTransactionDialog(self.db, tx, self).exec() == QDialog.DialogCode.Accepted:
            updated = self.db.get_transaction(ids[0])
            self.load_month(self.displayed_month, ids[0] if updated and updated["tx_date"][:7] == self.displayed_month else None)

    def delete_selected(self):
        ids = self.selected_ids()
        if not ids:
            QMessageBox.information(self, APP_NAME, "請先選取要刪除的資料。")
            return
        if QMessageBox.warning(
            self, APP_NAME, f"確定要刪除選取的 {len(ids)} 筆記帳資料嗎？\n刪除後無法復原。",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel,
            QMessageBox.StandardButton.Cancel,
        ) != QMessageBox.StandardButton.Yes:
            return
        try:
            self.db.delete_transactions(ids)
            self.load_month(self.displayed_month)
        except DatabaseError as e:
            QMessageBox.warning(self, APP_NAME, str(e))

    def confirm_lock(self):
        if self.current_data.get("missing_openings"):
            QMessageBox.warning(
                self, APP_NAME,
                "此月份尚有帳戶未設定期初餘額，請先完成設定再鎖定。",
            )
            return
        if QMessageBox.question(
            self, APP_NAME,
            f"確定 {self.displayed_month} 的記帳資料及期初餘額均已確認完成嗎？\n鎖定後，鎖定年月將更新為 {self.displayed_month}。",
        ) != QMessageBox.StandardButton.Yes:
            return
        self.db.set_locked_through(self.displayed_month)
        self.load_month(self.displayed_month)

    def export_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except Exception as e:
            QMessageBox.warning(self, APP_NAME, f"Excel 匯出元件載入失敗：{e}")
            return
        default = f"志遠記帳_{self.displayed_month.replace('/', '-')}.xlsx"
        export_dir = Path(self.config.get("export_directory") or Path.home()) if hasattr(self, "config") else Path.home()
        if not export_dir.exists():
            export_dir = Path.home()
        path, _ = QFileDialog.getSaveFileName(self, "單月匯出 Excel", str(export_dir / default), "Excel 活頁簿 (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            wb = Workbook(); ws = wb.active; ws.title = self.displayed_month.replace("/", "-")
            ws.merge_cells("A1:J1"); ws["A1"] = f"志遠記帳系統｜{self.displayed_month}"
            ws["A1"].font = Font(bold=True, size=16); ws["A1"].alignment = Alignment(horizontal="center")
            stats = [
                ("期初餘額", self.current_data["opening_total"]),
                ("收入小計", self.current_data["income_total"]),
                ("支出小計", self.current_data["expense_total"]),
                ("淨利" if self.current_data["net"] >= 0 else "淨損", abs(self.current_data["net"])),
                ("期末餘額", self.current_data["ending_total"]),
            ]
            for i, (label, value) in enumerate(stats, start=2):
                ws.cell(i, 1, label); ws.cell(i, 2, value); ws.cell(i, 2).number_format = '#,##0;[Red]-#,##0'
            if self.current_data["missing_openings"]:
                ws.cell(2, 4, "注意：此月份尚有帳戶未設定期初餘額")
                ws.cell(2, 4).font = Font(color="C62828", bold=True)
            header_row = 8
            for col, header in enumerate(LedgerTableModel.HEADERS, start=1):
                cell = ws.cell(header_row, col, header); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDE5EC"); cell.alignment = Alignment(horizontal="center")
            for rix, r in enumerate(self.current_data["rows"], start=header_row+1):
                values = [
                    r["tx_date"], r["account_name"], "收入" if r["kind"] == "income" else "支出",
                    r["category_name"] if r["kind"] == "income" else "",
                    r["summary"] if r["kind"] == "income" else "",
                    r["amount"] if r["kind"] == "income" else None,
                    r["category_name"] if r["kind"] == "expense" else "",
                    r["summary"] if r["kind"] == "expense" else "",
                    r["amount"] if r["kind"] == "expense" else None,
                    r["balance"],
                ]
                for cix, value in enumerate(values, start=1):
                    ws.cell(rix, cix, value)
                    if cix in (6, 9, 10) and value is not None:
                        ws.cell(rix, cix).number_format = '#,##0;[Red]-#,##0'
                fill = "FFFFFF" if (rix-header_row) % 2 == 1 else "F2F5F8"
                for cix in range(1, 11):
                    ws.cell(rix, cix).fill = PatternFill("solid", fgColor=fill)
            widths = [13, 14, 9, 14, 22, 14, 14, 22, 14, 16]
            for i, width in enumerate(widths, start=1):
                ws.column_dimensions[chr(64+i)].width = width
            ws.freeze_panes = f"A{header_row+1}"
            wb.save(path)
            QMessageBox.information(self, APP_NAME, f"Excel 匯出完成：\n{path}")
        except Exception as e:
            QMessageBox.warning(self, APP_NAME, f"Excel 匯出失敗：{e}")

    def refresh_current(self):
        self.load_month(self.displayed_month)


class MainWindow(QMainWindow):
    def __init__(self, db: Database, config: dict, startup_backup_error: str | None = None):
        startup_log("MainWindow.__init__ begin")
        super().__init__()
        self.db = db
        self.config = config
        self.startup_backup_error = startup_backup_error
        self.setWindowTitle(APP_NAME)
        self.setMinimumSize(1120, 760)
        self.resize(1120, 760)
        icon_path = app_root() / "app" / "resources" / "app.ico"
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))

        startup_log("MainWindow: reading latest transaction month")
        latest = db.latest_transaction_month() or date.today().strftime("%Y/%m")
        startup_log(f"MainWindow: latest month={latest}")
        startup_log("MainWindow: constructing InputTab")
        self.input_tab = InputTab(db, config, self.on_transaction_saved)
        startup_log("MainWindow: InputTab ready; constructing LedgerTab")
        self.ledger_tab = LedgerTab(db, latest, config)
        self.ledger_tab.restore_column_widths(config.get("ledger_column_widths"))
        startup_log("MainWindow: LedgerTab ready")

        self.tab_bar = QTabBar()
        self.tab_bar.setObjectName("mainTabBar")
        self.tab_bar.setExpanding(False)
        self.tab_bar.setDrawBase(False)
        self.tab_bar.addTab("輸入記帳")
        self.tab_bar.addTab("記帳資料表")

        self.pages = QStackedWidget()
        self.pages.addWidget(self.input_tab)
        self.pages.addWidget(self.ledger_tab)
        self.tab_bar.currentChanged.connect(self.pages.setCurrentIndex)
        self.tab_bar.currentChanged.connect(self.tab_changed)

        account_manage = no_tab_button("帳戶管理")
        account_manage.setObjectName("topActionButton")
        account_manage.clicked.connect(self.open_account_manager)
        category_manage = no_tab_button("收入支出科目管理")
        category_manage.setObjectName("topActionButton")
        category_manage.clicked.connect(self.open_category_manager)
        import_btn = no_tab_button("匯入")
        import_btn.setObjectName("topActionButton")
        import_btn.clicked.connect(self.open_import)
        settings = no_tab_button("⚙ 設定")
        settings.setObjectName("topActionButton")
        settings.clicked.connect(self.open_settings)

        central = QWidget()
        central_layout = QVBoxLayout(central)
        central_layout.setContentsMargins(16, 10, 16, 12)
        central_layout.setSpacing(0)

        top_nav = QHBoxLayout()
        top_nav.setContentsMargins(0, 0, 0, 0)
        top_nav.setSpacing(8)
        top_nav.addWidget(self.tab_bar, 0, Qt.AlignmentFlag.AlignBottom)
        top_nav.addStretch(1)
        top_nav.addWidget(account_manage, 0, Qt.AlignmentFlag.AlignTop)
        top_nav.addWidget(category_manage, 0, Qt.AlignmentFlag.AlignTop)
        top_nav.addWidget(import_btn, 0, Qt.AlignmentFlag.AlignTop)
        top_nav.addWidget(settings, 0, Qt.AlignmentFlag.AlignTop)
        central_layout.addLayout(top_nav)

        page_container = QFrame()
        page_container.setObjectName("pageContainer")
        page_layout = QVBoxLayout(page_container)
        page_layout.setContentsMargins(0, 0, 0, 0)
        page_layout.addWidget(self.pages)
        central_layout.addWidget(page_container, 1)
        self.setCentralWidget(central)
        self._restore_window_state()
        QTimer.singleShot(0, self.input_tab.date_edit.setFocus)
        if startup_backup_error:
            QTimer.singleShot(700, lambda: QMessageBox.warning(self, APP_NAME, startup_backup_error))
        startup_log("MainWindow.__init__ complete")

    def _restore_window_state(self):
        geom = self.config.get("window_geometry")
        if isinstance(geom, dict):
            try:
                rect = QRect(int(geom["x"]), int(geom["y"]), int(geom["w"]), int(geom["h"]))
                visible = any(rect.intersects(screen.availableGeometry()) for screen in QApplication.screens())
                if visible:
                    self.setGeometry(rect)
            except Exception:
                pass
        if self.config.get("window_maximized"):
            QTimer.singleShot(0, self.showMaximized)

    def _save_window_state(self):
        rect = self.normalGeometry() if self.isMaximized() else self.geometry()
        self.config["window_geometry"] = {"x": rect.x(), "y": rect.y(), "w": rect.width(), "h": rect.height()}
        self.config["window_maximized"] = bool(self.isMaximized())
        self.config["ledger_column_widths"] = self.ledger_tab.column_widths()
        save_config(self.config)

    def on_transaction_saved(self, month: str):
        if self.ledger_tab.displayed_month == month:
            self.ledger_tab.refresh_current()

    def tab_changed(self, index: int):
        if index == 1:
            self.ledger_tab.refresh_current()
            position = self.config.get("ledger_entry_position", "latest")
            QTimer.singleShot(0, lambda: self.ledger_tab.scroll_to_entry_position(position))

    def open_account_manager(self):
        self.input_tab.open_account_manager()
        self.ledger_tab.refresh_current()

    def open_category_manager(self):
        self.input_tab.open_category_manager()
        self.ledger_tab.refresh_current()

    def open_import(self):
        dlg = ImportTransactionsDialog(self.db, self.config, self)
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.imported_count > 0:
            save_config(self.config)
            self.input_tab.refresh_lists()
            self.ledger_tab.refresh_current()

    def open_settings(self):
        dlg = SettingsDialog(
            self.db, self.config, self.change_database_location, self.restore_database,
            self.clear_data, self.settings_changed, self,
        )
        dlg.exec()
        save_config(self.config)
        self.input_tab.refresh_lists()
        self.ledger_tab.refresh_current()

    def settings_changed(self):
        self.input_tab.refresh_lists()
        self.ledger_tab.refresh_current()

    def clear_data(self):
        if background_tasks_running():
            raise DatabaseError("Google Drive 正在同步，完成後才能重置帳本。")
        restore_dir = self.db.path.parent / "RestoreBackup"
        before = restore_dir / f"CYacc_beforeclear_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.db"
        self.db.backup_to(before)
        old_config = dict(self.config)
        try:
            self.db.reset_local_ledger()
            # Keep the active DB location so the next launch cannot open an old
            # default database by accident. Everything else, including Drive
            # authorization and UI preferences, returns to its initial state.
            self.config.clear()
            self.config["database_path"] = str(self.db.path)
            save_config(self.config)
            safe_copy(config_path(), config_path().with_suffix(".json.bak"))
        except Exception:
            self.config.clear()
            self.config.update(old_config)
            try:
                save_config(self.config)
            except Exception as restore_config_error:
                error_log(f"failed to restore configuration after clear: {restore_config_error}", sys.exc_info())
            restored, reason = self.restore_database(before)
            if not restored:
                error_log(f"failed to restore database after clear: {reason}")
                raise DatabaseError(f"重置未完成；還原原帳本失敗：{reason}。請保留 {before}。")
            raise
        month = date.today().strftime("%Y/%m")
        self.ledger_tab.month_spin.set_month(month)
        self.ledger_tab.load_month(month)
        self.input_tab.refresh_lists()

    def change_database_location(self, folder: Path) -> tuple[bool, str]:
        old_path = self.db.path
        try:
            folder.mkdir(parents=True, exist_ok=True)
            target = folder / DB_FILENAME
            if target.exists():
                ok, msg = self.db.validate_database_file(target)
                if not ok:
                    return False, f"目標資料夾中的資料庫無效：{msg}"
                if QMessageBox.question(self, APP_NAME, "目標資料夾已有資料庫，是否改用該資料庫？") != QMessageBox.StandardButton.Yes:
                    return False, "已取消切換"
            else:
                self.db.backup_to(target)
            self.db.close(); self.db.path = target; self.db.open()
            ok, msg = self.db.integrity_check(True)
            if not ok:
                raise DatabaseError(msg)
            self.config["database_path"] = str(target)
            save_config(self.config)
            app_log(f"database location changed to {target}")
            return True, ""
        except Exception as e:
            error_log(f"change database location failed: {e}", sys.exc_info())
            try:
                self.db.close(); self.db.path = old_path; self.db.open()
            except Exception:
                pass
            return False, str(e)

    def restore_database(self, source: Path) -> tuple[bool, str]:
        ok, msg = self.db.validate_database_file(source)
        if not ok:
            return False, msg
        restore_dir = self.db.path.parent / "RestoreBackup"
        restore_dir.mkdir(parents=True, exist_ok=True)
        before = restore_dir / f"CYacc_restorebefore_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
        current = self.db.path
        candidate = current.with_suffix(".restore_candidate.tmp")
        rollback = current.with_suffix(".restore_original.tmp")
        try:
            self.db.backup_to(before)
            shutil.copy2(source, candidate)
            ok, msg = self.db.validate_database_file(candidate)
            if not ok:
                raise DatabaseError(f"還原檔驗證失敗：{msg}")
            self.db.checkpoint()
            self.db.close()
            for sidecar in (Path(str(current)+"-wal"), Path(str(current)+"-shm")):
                sidecar.unlink(missing_ok=True)
            rollback.unlink(missing_ok=True)
            if current.exists():
                os.replace(current, rollback)
            os.replace(candidate, current)
            self.db.open()
            ok, msg = self.db.integrity_check(True)
            if not ok:
                raise DatabaseError(msg)
            rollback.unlink(missing_ok=True)
            app_log(f"database restored from {source}; pre-restore backup={before}")
            return True, ""
        except Exception as e:
            error_log(f"restore database failed: {e}", sys.exc_info())
            try:
                self.db.close()
                candidate.unlink(missing_ok=True)
                if rollback.exists():
                    current.unlink(missing_ok=True)
                    os.replace(rollback, current)
                elif before.exists():
                    shutil.copy2(before, current)
                self.db.open()
            except Exception as rollback_error:
                error_log(f"restore rollback failed: {rollback_error}", sys.exc_info())
            return False, str(e)

    def closeEvent(self, event: QCloseEvent):  # noqa: N802
        if background_tasks_running():
            QMessageBox.information(self, APP_NAME, "Google Drive 正在背景同步，完成後即可安全關閉程式。")
            event.ignore()
            return
        if self.input_tab.has_unsaved_content():
            answer = QMessageBox.question(
                self, APP_NAME,
                "目前有尚未存入的記帳內容，確定要關閉程式嗎？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No,
            )
            if answer != QMessageBox.StandardButton.Yes:
                event.ignore(); return
        self._save_window_state()
        self.db.close()
        event.accept()




def install_exception_hook():
    def hook(exc_type, exc_value, exc_tb):
        startup_log(f"RUNTIME EXCEPTION: {exc_type.__name__}: {exc_value}")
        try:
            traceback.print_exception(exc_type, exc_value, exc_tb, file=_STARTUP_LOG_FILE)
            _STARTUP_LOG_FILE.flush()
            os.fsync(_STARTUP_LOG_FILE.fileno())
        except Exception:
            pass
        error_log(f"UNHANDLED RUNTIME EXCEPTION: {exc_type.__name__}: {exc_value}", (exc_type, exc_value, exc_tb))
        app = QApplication.instance()
        if app is not None:
            try:
                QMessageBox.critical(None, APP_NAME, "程式發生未預期錯誤，詳細內容已寫入 Data\\startup_trace.log。")
            except Exception:
                pass
    sys.excepthook = hook
    startup_log("runtime exception hook installed")

def run_auto_backup(db: Database, config: dict) -> str | None:
    today = date.today()
    last_text = config.get("last_auto_backup_date", "")
    due = True
    if last_text:
        try:
            last = datetime.strptime(last_text, "%Y/%m/%d").date()
            due = (today - last).days >= 3
        except Exception:
            due = True
    if not due:
        return None
    backup_dir = db.path.parent / "AutoBackup"
    backup_dir.mkdir(parents=True, exist_ok=True)
    target = backup_dir / f"{BACKUP_PREFIX}{today.strftime('%Y%m%d')}.db"
    try:
        if not target.exists():
            db.backup_to(target)
        else:
            ok, msg = db.validate_database_file(target)
            if not ok:
                target.unlink(missing_ok=True)
                db.backup_to(target)
        config["last_auto_backup_date"] = today.strftime("%Y/%m/%d")
        save_config(config)
        backups = sorted(backup_dir.glob(f"{BACKUP_PREFIX}*.db"), key=lambda p: p.name, reverse=True)
        for old in backups[30:]:
            try:
                old.unlink()
            except Exception:
                pass
        return None
    except Exception:
        return "自動備份失敗，請確認資料庫資料夾是否有寫入權限。"


def run_google_drive_sync(db: Database, config: dict) -> str | None:
    if not bool(config.get("google_drive_sync_enabled", False)):
        return None
    client = GoogleDriveClient(config)
    if not client.is_connected:
        return "Google Drive 尚未完成授權，已略過雲端同步。"
    backup_dir = db.path.parent / "AutoBackup"
    latest = newest_local_backup(backup_dir)
    if latest is None:
        return None
    if config.get("google_drive_last_backup_name") == latest.name:
        return None
    try:
        updated = upload_backup_and_cleanup(latest, config, 30)
        for key, value in updated.items():
            if key.startswith("google_"):
                config[key] = value
        save_config(config)
        return None
    except Exception as exc:
        return f"Google Drive 自動同步失敗：{exc}"


def main() -> int:
    startup_log("main() begin")
    sync_version_marker()
    app_log(f"application launch requested, version={APP_VERSION}")
    if not acquire_single_instance():
        return 0
    try_set_zh_tw_locale()
    startup_log("system collation locale attempt complete")

    # Qt 6 on Windows otherwise follows the OS dark theme for native title bars
    # and the unstyled item views.  Keep the existing light application theme.
    if sys.platform == "win32" and "QT_QPA_PLATFORM" not in os.environ:
        os.environ["QT_QPA_PLATFORM"] = "windows:darkmode=0"
    startup_log("creating QApplication")
    app = QApplication(sys.argv)
    app.styleHints().setColorScheme(Qt.ColorScheme.Light)
    startup_log("QApplication created")
    install_exception_hook()
    app.setApplicationName(APP_NAME)
    app.setApplicationVersion(APP_VERSION)
    app.setOrganizationName("C.C.LIU")
    app.setStyleSheet(APP_STYLE)
    app._chinese_button_filter = ChineseStandardButtonFilter(app)
    app.installEventFilter(app._chinese_button_filter)
    icon_path = app_root() / "app" / "resources" / "app.ico"
    if icon_path.exists():
        app.setWindowIcon(QIcon(str(icon_path)))
    startup_log("application metadata and style applied")

    startup_log("loading configuration")
    config_existed_at_start = config_path().exists()
    try:
        config, config_warning = load_config_with_status()
    except ConfigError as exc:
        startup_log(f"configuration load failed safely: {exc}")
        error_log(f"configuration load failed safely: {exc}", sys.exc_info())
        QMessageBox.critical(
            None, APP_NAME,
            "設定檔已損壞，程式已停止開啟，以免誤用新的空白帳本。\n\n"
            f"原因：{exc}\n\n請保留 Data 資料夾並提供 startup_trace.log 協助處理。"
        )
        return 2
    configured_database = str(config.get("database_path") or "").strip()
    db_path = Path(configured_database or default_database_path())
    try:
        database_missing_or_empty = not db_path.exists() or db_path.stat().st_size <= 0
    except OSError:
        database_missing_or_empty = True
    if (configured_database or config_existed_at_start) and database_missing_or_empty:
        startup_log(f"configured database missing or empty; startup stopped: {db_path}")
        QMessageBox.critical(
            None, APP_NAME,
            "原本使用的資料庫不存在或是空白檔案，程式已停止開啟，且不會自動建立新帳本。\n\n"
            f"原資料庫位置：{db_path}\n\n請確認外接磁碟／網路資料夾是否已連線，或從備份還原。"
        )
        return 2
    startup_log(f"opening database: {db_path}")
    try:
        db = Database(db_path)
    except Exception as exc:
        startup_log(f"database open failed: {exc}")
        error_log(f"database open failed: {exc}", sys.exc_info())
        ready_path = _STARTUP_DATA_DIR / "startup_ready.flag"
        try:
            ready_path.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
        except Exception:
            pass
        backup_folder = db_path.parent / "AutoBackup"
        QMessageBox.critical(
            None, APP_NAME,
            f"資料庫無法開啟，程式已停止寫入以保護資料。\n\n資料庫：{db_path}\n原因：{exc}\n\n可檢查最近備份：{backup_folder}"
        )
        return 2
    startup_log("database open and schema initialization complete")
    app_log(f"database integrity verified: {db_path}")
    if config.get("database_path") != str(db_path):
        config["database_path"] = str(db_path)
        try:
            save_config(config)
        except Exception as exc:
            config_warning = (config_warning + "\n" if config_warning else "") + f"無法保存資料庫位置安全設定：{exc}"

    startup_log("constructing main window")
    win = MainWindow(db, config, None)
    startup_log("main window constructed")
    if config_warning:
        QTimer.singleShot(700, lambda message=config_warning: QMessageBox.warning(win, APP_NAME, message))

    def force_window_visible(stage: str) -> None:
        """Show the Qt window and, on Windows, explicitly restore the native HWND.

        Windows honors STARTUPINFO.wShowWindow on the first ShowWindow call. A launcher
        that supplied SW_HIDE can therefore make Qt report visible=True while the native
        window remains invisible. The launcher no longer supplies SW_HIDE, and this
        second native show is retained as a defensive measure.
        """
        try:
            should_maximize = bool(config.get("window_maximized")) or win.isMaximized()
            if should_maximize:
                win.showMaximized()
            else:
                win.show()
            win.raise_()
            win.activateWindow()
            app.processEvents()
            native_detail = "native=n/a"
            if sys.platform == "win32":
                import ctypes

                hwnd = int(win.winId())
                user32 = ctypes.windll.user32
                before = int(user32.IsWindowVisible(hwnd))
                SW_RESTORE = 9
                SW_SHOWMAXIMIZED = 3
                HWND_TOPMOST = -1
                HWND_NOTOPMOST = -2
                SWP_NOSIZE = 0x0001
                SWP_NOMOVE = 0x0002
                SWP_SHOWWINDOW = 0x0040
                flags = SWP_NOSIZE | SWP_NOMOVE | SWP_SHOWWINDOW
                user32.ShowWindow(hwnd, SW_SHOWMAXIMIZED if should_maximize else SW_RESTORE)
                user32.SetWindowPos(hwnd, HWND_TOPMOST, 0, 0, 0, 0, flags)
                user32.SetWindowPos(hwnd, HWND_NOTOPMOST, 0, 0, 0, 0, flags)
                user32.BringWindowToTop(hwnd)
                user32.SetForegroundWindow(hwnd)
                after = int(user32.IsWindowVisible(hwnd))
                native_detail = f"hwnd={hwnd} native_visible_before={before} native_visible_after={after}"
            startup_log(
                f"force_window_visible[{stage}] | qt_visible={win.isVisible()} | {native_detail} | "
                f"geometry={win.geometry().x()},{win.geometry().y()},"
                f"{win.geometry().width()}x{win.geometry().height()}"
            )
        except Exception as exc:
            startup_log(f"force_window_visible[{stage}] failed: {exc}")

    ready_path = _STARTUP_DATA_DIR / "startup_ready.flag"

    def finalize_startup() -> None:
        force_window_visible("event-loop-start")
        try:
            ready_path.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
            startup_log(f"startup ready flag written: {ready_path}")
        except Exception as exc:
            startup_log(f"could not write startup ready flag: {exc}")
        try:
            faulthandler.cancel_dump_traceback_later()
            startup_log("startup watchdog cancelled: UI reached ready state")
        except Exception as exc:
            startup_log(f"could not cancel startup watchdog: {exc}")

    # The first call creates the native window. Subsequent queued calls ensure Windows
    # cannot leave it hidden because of inherited startup-show state.
    force_window_visible("before-event-loop")
    QTimer.singleShot(0, finalize_startup)
    QTimer.singleShot(350, lambda: force_window_visible("350ms"))
    QTimer.singleShot(1200, lambda: force_window_visible("1200ms"))

    def deferred_auto_backup():
        startup_log("deferred automatic backup check begin")
        error = run_auto_backup(db, config)
        if error:
            startup_log(f"automatic backup failed: {error}")
            QMessageBox.warning(win, APP_NAME, error)
            return
        if not bool(config.get("google_drive_sync_enabled", False)):
            startup_log("deferred automatic backup check complete; cloud sync disabled")
            return
        backup_dir = db.path.parent / "AutoBackup"
        latest = newest_local_backup(backup_dir)
        if latest is None or config.get("google_drive_last_backup_name") == latest.name:
            startup_log("deferred automatic backup check complete; no new cloud backup")
            return
        snapshot = dict(config)

        def cloud_task():
            return upload_backup_and_cleanup(latest, snapshot, 30)

        def cloud_finished(updated, exc):
            if exc is not None:
                message = f"Google Drive 自動同步失敗：{exc}"
                startup_log(message)
                QMessageBox.warning(win, APP_NAME, message)
            elif isinstance(updated, dict):
                for key, value in updated.items():
                    if key.startswith("google_"):
                        config[key] = value
                save_config(config)
                startup_log("Google Drive automatic sync completed in background")
            startup_log("deferred automatic backup check complete")

        win._background_drive_sync = start_background_task(cloud_task, cloud_finished)

    win._deferred_auto_backup = deferred_auto_backup
    QTimer.singleShot(800, deferred_auto_backup)
    QTimer.singleShot(10, lambda: startup_log("Qt event loop processed its first queued callback"))
    startup_log("entering Qt event loop")
    exit_code = app.exec()
    startup_log(f"Qt event loop exited with code {exit_code}")
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
