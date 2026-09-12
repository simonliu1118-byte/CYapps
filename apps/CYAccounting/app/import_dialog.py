from __future__ import annotations

import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from PySide6.QtCore import Qt
from PySide6.QtGui import QIntValidator
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDialog,
    QFileDialog,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QPushButton,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from db import Database, DatabaseError
from gdrive import DriveFile, GoogleDriveClient, GoogleDriveError, GOOGLE_SHEET_MIME, XLS_MIME, XLSM_MIME
from util import APP_NAME, MAX_AMOUNT, format_date, is_month_locked, month_key_from_date, normalize_date_input, parse_date, weighted_units


class DriveFileDialog(QDialog):
    """Small native Drive browser for spreadsheet files visible to the authorized account."""

    def __init__(self, client: GoogleDriveClient, parent=None):
        super().__init__(parent)
        self.client = client
        self.selected_file: DriveFile | None = None
        self.setWindowTitle("從 Google Drive 選擇帳簿")
        self.resize(680, 480)
        outer = QVBoxLayout(self)

        row = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText("搜尋 Google Drive 檔名")
        self.search.returnPressed.connect(self.refresh)
        search_btn = QPushButton("搜尋")
        search_btn.clicked.connect(self.refresh)
        refresh_btn = QPushButton("重新整理")
        refresh_btn.clicked.connect(self.refresh)
        row.addWidget(self.search, 1)
        row.addWidget(search_btn)
        row.addWidget(refresh_btn)
        outer.addLayout(row)

        self.list = QListWidget()
        self.list.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.list.itemDoubleClicked.connect(lambda _item: self.accept_selected())
        outer.addWidget(self.list, 1)

        hint = QLabel("可選擇 .xlsx、.xlsm 或 Google 試算表；舊式 .xls 請先另存為 .xlsx。")
        hint.setStyleSheet("color:#667085;")
        outer.addWidget(hint)

        buttons = QHBoxLayout()
        buttons.addStretch(1)
        cancel = QPushButton("取消"); cancel.clicked.connect(self.reject)
        choose = QPushButton("選擇"); choose.clicked.connect(self.accept_selected)
        buttons.addWidget(cancel); buttons.addWidget(choose)
        outer.addLayout(buttons)
        self.refresh()

    def refresh(self):
        self.list.clear()
        try:
            files = self.client.list_spreadsheet_files(self.search.text())
        except GoogleDriveError as exc:
            QMessageBox.warning(self, APP_NAME, str(exc))
            return
        for file in files:
            modified = file.modified_time.replace("T", " ")[:16] if file.modified_time else ""
            item = QListWidgetItem(f"{file.name}\n{file.type_label}　{modified}")
            item.setData(Qt.ItemDataRole.UserRole, file)
            if file.mime_type == XLS_MIME:
                item.setToolTip("舊式 .xls 可選取，但匯入前仍需另存為 .xlsx。")
            self.list.addItem(item)
        if self.list.count() == 0:
            item = QListWidgetItem("找不到可匯入的試算表")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.list.addItem(item)

    def accept_selected(self):
        item = self.list.currentItem()
        if item is None:
            return
        file = item.data(Qt.ItemDataRole.UserRole)
        if not isinstance(file, DriveFile):
            return
        self.selected_file = file
        self.accept()


def _normalize_header(text: Any) -> str:
    s = str(text or "").strip().lower()
    s = re.sub(r"[\s_\-－—｜|:：()（）\[\]【】]", "", s)
    return s


def _excel_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return format_date(value.date())
    if isinstance(value, date):
        return format_date(value)
    text = str(value).strip()
    if not text:
        return None
    # Format-only transformation remains separate from validity checking.
    text = normalize_date_input(text)
    d = parse_date(text)
    if d:
        return format_date(d)
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            d = datetime.strptime(text, fmt).date()
            return format_date(d)
        except Exception:
            pass
    # Windows/Python does not consistently support %-d, so parse variable-width
    # slash/dash dates explicitly.
    m = re.fullmatch(r"(\d{4})[/.\-](\d{1,2})[/.\-](\d{1,2})", text)
    if m:
        try:
            return format_date(date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
        except Exception:
            return None
    return None


def _excel_amount(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = str(value).strip().replace(",", "").replace("$", "").replace("＄", "")
    if not text:
        return None
    try:
        number = float(text)
        return int(number) if number.is_integer() else None
    except Exception:
        return None


def _kind_value(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    if text in {"收入", "收", "income", "in", "+"}:
        return "income"
    if text in {"支出", "支", "expense", "out", "-"}:
        return "expense"
    return None


class ImportTransactionsDialog(QDialog):
    """Import external Excel/Google Sheets data with explicit column mapping."""

    MODE_SPLIT = "split"
    MODE_UNIFIED = "unified"

    SPLIT_FIELDS = [
        ("date", "日期", True),
        ("account", "帳戶", True),
        ("income_category", "收入科目", False),
        ("income_summary", "收入摘要", False),
        ("income_amount", "收入金額", False),
        ("expense_category", "支出科目", False),
        ("expense_summary", "支出摘要", False),
        ("expense_amount", "支出金額", False),
    ]
    UNIFIED_FIELDS = [
        ("date", "日期", True),
        ("account", "帳戶", True),
        ("kind", "收支", True),
        ("category", "科目", True),
        ("summary", "摘要", False),
        ("amount", "金額", True),
    ]

    SYNONYMS = {
        "date": {"日期", "date", "交易日期", "記帳日期"},
        "account": {"帳戶", "account", "帳號", "付款帳戶", "收款帳戶"},
        "kind": {"收支", "類型", "收入支出", "kind", "type"},
        "category": {"科目", "類別", "分類", "category"},
        "summary": {"摘要", "備註", "說明", "memo", "remark", "description"},
        "amount": {"金額", "amount"},
        "income_category": {"收入科目", "收入類別", "收入分類"},
        "income_summary": {"收入摘要", "收入備註", "收入說明"},
        "income_amount": {"收入金額", "收入", "收入額"},
        "expense_category": {"支出科目", "支出類別", "支出分類"},
        "expense_summary": {"支出摘要", "支出備註", "支出說明"},
        "expense_amount": {"支出金額", "支出", "支出額"},
    }

    def __init__(self, db: Database, config: dict, parent=None):
        super().__init__(parent)
        self.db = db
        self.config = config
        self.workbook = None
        self.current_path: Path | None = None
        self.temp_path: Path | None = None
        self.imported_count = 0
        self.setWindowTitle("匯入外部帳簿")
        self.resize(900, 690)
        self.setObjectName("importDialog")
        self.setStyleSheet(
            "QDialog#importDialog QPushButton, QDialog#importDialog QLineEdit, QDialog#importDialog QComboBox {"
            "min-height:24px; max-height:27px; padding-top:0px; padding-bottom:0px;}"
        )
        self._build()

    def _build(self):
        outer = QVBoxLayout(self)
        outer.setSpacing(8)

        source = QGroupBox("匯入來源")
        sl = QGridLayout(source)
        local = QPushButton("從本機選擇")
        local.clicked.connect(self.choose_local)
        drive = QPushButton("從 Google Drive 選擇")
        drive.clicked.connect(self.choose_drive)
        self.file_label = QLabel("尚未選擇檔案")
        self.file_label.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        sl.addWidget(local, 0, 0)
        sl.addWidget(drive, 0, 1)
        sl.addWidget(self.file_label, 0, 2, 1, 3)
        sl.addWidget(QLabel("工作表："), 1, 0)
        self.sheet_combo = QComboBox(); self.sheet_combo.currentIndexChanged.connect(self.reload_headers)
        sl.addWidget(self.sheet_combo, 1, 1)
        sl.addWidget(QLabel("標題列："), 1, 2)
        self.header_row = QLineEdit("1")
        self.header_row.setValidator(QIntValidator(1, 9999, self.header_row))
        self.header_row.setFixedWidth(55)
        self.header_row.editingFinished.connect(self.reload_headers)
        sl.addWidget(self.header_row, 1, 3)
        sl.addWidget(QLabel("格式："), 1, 4)
        self.mode_combo = QComboBox()
        self.mode_combo.addItem("收入／支出分欄", self.MODE_SPLIT)
        self.mode_combo.addItem("單一收支欄", self.MODE_UNIFIED)
        self.mode_combo.currentIndexChanged.connect(self._mode_changed)
        sl.addWidget(self.mode_combo, 1, 5)
        sl.setColumnStretch(2, 1)
        outer.addWidget(source)

        mapping_box = QGroupBox("欄位對應")
        mv = QVBoxLayout(mapping_box)
        self.mapping_stack = QStackedWidget()
        self.split_page, self.split_mapping = self._make_mapping_page(self.SPLIT_FIELDS)
        self.unified_page, self.unified_mapping = self._make_mapping_page(self.UNIFIED_FIELDS)
        self.mapping_stack.addWidget(self.split_page)
        self.mapping_stack.addWidget(self.unified_page)
        mv.addWidget(self.mapping_stack)
        note = QLabel("不存在於目前主檔的帳戶／科目會以歷史文字匯入，不會自動新增或改動帳戶、科目設定。")
        note.setStyleSheet("color:#667085;")
        mv.addWidget(note)
        outer.addWidget(mapping_box)

        preview_box = QGroupBox("匯入預覽（前 20 筆）")
        pv = QVBoxLayout(preview_box)
        self.preview = QTableWidget(0, 7)
        self.preview.setHorizontalHeaderLabels(["Excel列", "日期", "帳戶", "收支", "科目", "摘要", "金額／狀態"])
        self.preview.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.preview.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.preview.verticalHeader().setVisible(False)
        hh = self.preview.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        hh.setSectionResizeMode(5, QHeaderView.ResizeMode.Stretch)
        pv.addWidget(self.preview)
        action = QHBoxLayout()
        self.skip_duplicates = QCheckBox("略過完全相同的重複資料")
        self.skip_duplicates.setChecked(True)
        action.addWidget(self.skip_duplicates)
        action.addStretch(1)
        refresh = QPushButton("重新分析"); refresh.clicked.connect(self.refresh_preview)
        action.addWidget(refresh)
        pv.addLayout(action)
        outer.addWidget(preview_box, 1)

        bottom = QHBoxLayout()
        bottom.addStretch(1)
        cancel = QPushButton("取消"); cancel.clicked.connect(self.reject)
        self.import_btn = QPushButton("開始匯入"); self.import_btn.clicked.connect(self.do_import)
        self.import_btn.setEnabled(False)
        bottom.addWidget(cancel); bottom.addWidget(self.import_btn)
        outer.addLayout(bottom)

    def _make_mapping_page(self, fields: list[tuple[str, str, bool]]) -> tuple[QWidget, dict[str, QComboBox]]:
        page = QWidget()
        grid = QGridLayout(page)
        grid.setContentsMargins(0, 0, 0, 0)
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(5)
        mapping: dict[str, QComboBox] = {}
        for index, (key, label, required) in enumerate(fields):
            row = index // 4
            col = (index % 4) * 2
            lab = QLabel(label + (" *" if required else ""))
            combo = QComboBox()
            combo.setMinimumWidth(130)
            combo.currentIndexChanged.connect(self.refresh_preview)
            grid.addWidget(lab, row, col)
            grid.addWidget(combo, row, col + 1)
            mapping[key] = combo
        grid.setColumnStretch(7, 1)
        return page, mapping

    def _mode_changed(self):
        mode = self.mode_combo.currentData()
        self.mapping_stack.setCurrentIndex(0 if mode == self.MODE_SPLIT else 1)
        self.auto_map()
        self.refresh_preview()

    def _current_mapping(self) -> dict[str, QComboBox]:
        return self.split_mapping if self.mode_combo.currentData() == self.MODE_SPLIT else self.unified_mapping

    def choose_local(self):
        initial = self.config.get("import_directory") or str(Path.home())
        path, _ = QFileDialog.getOpenFileName(
            self, "選擇 Excel 帳簿", initial, "Excel 活頁簿 (*.xlsx *.xlsm *.xls)"
        )
        if not path:
            return
        self.config["import_directory"] = str(Path(path).parent)
        self.load_file(Path(path), Path(path).name)

    def choose_drive(self):
        client = GoogleDriveClient(self.config)
        if not client.is_connected:
            QMessageBox.information(self, APP_NAME, "請先到「設定 → Google Drive」完成授權。")
            return
        dlg = DriveFileDialog(client, self)
        if dlg.exec() != QDialog.DialogCode.Accepted or dlg.selected_file is None:
            return
        file = dlg.selected_file
        if file.mime_type == XLS_MIME:
            QMessageBox.warning(self, APP_NAME, "舊式 .xls 目前無法直接解析，請在 Excel 或 Google 試算表另存為 .xlsx 後再匯入。")
            return
        suffix = ".xlsm" if file.mime_type == XLSM_MIME or file.name.lower().endswith(".xlsm") else ".xlsx"
        tmp_dir = Path(tempfile.gettempdir()) / "CYAccountingImport"
        tmp_dir.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix="cy_import_", suffix=suffix, dir=str(tmp_dir))
        try:
            import os
            os.close(fd)
        except Exception:
            pass
        try:
            self.temp_path = client.download_spreadsheet(file, Path(name))
            self.load_file(self.temp_path, f"Google Drive：{file.name}")
        except GoogleDriveError as exc:
            QMessageBox.warning(self, APP_NAME, str(exc))

    def load_file(self, path: Path, display_name: str):
        if path.suffix.lower() == ".xls":
            QMessageBox.warning(self, APP_NAME, "舊式 .xls 目前無法直接解析，請另存為 .xlsx 後再匯入。")
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"無法讀取 Excel：{exc}")
            return
        if self.workbook is not None:
            try:
                self.workbook.close()
            except Exception:
                pass
        self.workbook = wb
        self.current_path = path
        self.file_label.setText(display_name)
        self.sheet_combo.blockSignals(True)
        self.sheet_combo.clear(); self.sheet_combo.addItems(wb.sheetnames)
        self.sheet_combo.blockSignals(False)
        detected = self._detect_header_row()
        self.header_row.setText(str(detected))
        self.reload_headers()

    def _detect_header_row(self) -> int:
        ws = self._worksheet()
        if ws is None:
            return 1
        date_names = {_normalize_header(x) for x in self.SYNONYMS["date"]}
        account_names = {_normalize_header(x) for x in self.SYNONYMS["account"]}
        amount_names = {_normalize_header(x) for x in (self.SYNONYMS["amount"] | self.SYNONYMS["income_amount"] | self.SYNONYMS["expense_amount"])}
        try:
            for row_no, values in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
                normalized = {_normalize_header(v) for v in values if v not in (None, "")}
                if normalized & date_names and normalized & account_names and normalized & amount_names:
                    return row_no
        except Exception:
            pass
        return 1

    def _worksheet(self):
        if self.workbook is None or not self.sheet_combo.currentText():
            return None
        return self.workbook[self.sheet_combo.currentText()]

    def _header_row_number(self) -> int:
        try:
            return max(1, int(self.header_row.text()))
        except Exception:
            return 1

    def reload_headers(self):
        ws = self._worksheet()
        if ws is None:
            self.import_btn.setEnabled(False)
            return
        row_no = self._header_row_number()
        try:
            values = next(ws.iter_rows(min_row=row_no, max_row=row_no, values_only=True))
        except StopIteration:
            values = []
        headers = []
        for idx, value in enumerate(values, start=1):
            text = str(value or "").strip()
            headers.append(text or f"第{idx}欄")
        for mapping in (self.split_mapping, self.unified_mapping):
            for combo in mapping.values():
                combo.blockSignals(True)
                combo.clear()
                combo.addItem("（未指定）", -1)
                for idx, text in enumerate(headers):
                    combo.addItem(text, idx)
                combo.blockSignals(False)
        normalized_headers = {_normalize_header(h) for h in headers}
        has_split = bool(
            normalized_headers & {_normalize_header(x) for x in (self.SYNONYMS["income_amount"] | self.SYNONYMS["expense_amount"])}
        )
        has_unified = (
            bool(normalized_headers & {_normalize_header(x) for x in self.SYNONYMS["kind"]})
            and bool(normalized_headers & {_normalize_header(x) for x in self.SYNONYMS["amount"]})
        )
        target_mode = self.MODE_SPLIT if has_split or not has_unified else self.MODE_UNIFIED
        mode_index = self.mode_combo.findData(target_mode)
        if mode_index >= 0:
            self.mode_combo.blockSignals(True)
            self.mode_combo.setCurrentIndex(mode_index)
            self.mapping_stack.setCurrentIndex(0 if target_mode == self.MODE_SPLIT else 1)
            self.mode_combo.blockSignals(False)
        self.auto_map()
        self.refresh_preview()
        self.import_btn.setEnabled(bool(headers))

    def auto_map(self):
        mapping = self._current_mapping()
        headers: list[tuple[int, str]] = []
        if mapping:
            sample = next(iter(mapping.values()))
            for i in range(1, sample.count()):
                headers.append((int(sample.itemData(i)), _normalize_header(sample.itemText(i))))
        for key, combo in mapping.items():
            targets = {_normalize_header(x) for x in self.SYNONYMS.get(key, set())}
            found = -1
            for column_index, normalized in headers:
                if normalized in targets:
                    found = column_index
                    break
            combo.blockSignals(True)
            idx = combo.findData(found) if found >= 0 else 0
            combo.setCurrentIndex(idx if idx >= 0 else 0)
            combo.blockSignals(False)

    def _mapped_value(self, mapping: dict[str, QComboBox], key: str, values: tuple) -> Any:
        combo = mapping.get(key)
        if combo is None:
            return None
        index = int(combo.currentData() if combo.currentData() is not None else -1)
        return values[index] if 0 <= index < len(values) else None

    def _normalized_row(self, excel_row: int, values: tuple) -> tuple[dict | None, str | None]:
        mode = self.mode_combo.currentData()
        mapping = self._current_mapping()
        tx_date = _excel_date(self._mapped_value(mapping, "date", values))
        account = str(self._mapped_value(mapping, "account", values) or "").strip()
        if not tx_date and not account and all(v in (None, "") for v in values):
            return None, None
        if not tx_date:
            return None, "日期錯誤"
        if not account:
            return None, "帳戶空白"
        if is_month_locked(month_key_from_date(parse_date(tx_date)), self.db.locked_through()):
            return None, f"{tx_date[:7]} 已鎖定"

        if mode == self.MODE_UNIFIED:
            kind = _kind_value(self._mapped_value(mapping, "kind", values))
            category = str(self._mapped_value(mapping, "category", values) or "").strip()
            summary = str(self._mapped_value(mapping, "summary", values) or "").strip()
            amount = _excel_amount(self._mapped_value(mapping, "amount", values))
            if kind is None:
                return None, "收支必須為收入或支出"
            if not category:
                return None, "科目空白"
            if amount is None or amount < 1:
                return None, "金額錯誤"
            if amount > MAX_AMOUNT:
                return None, "金額最多7位數"
        else:
            income_amount = _excel_amount(self._mapped_value(mapping, "income_amount", values))
            expense_amount = _excel_amount(self._mapped_value(mapping, "expense_amount", values))
            has_income = income_amount is not None and income_amount != 0
            has_expense = expense_amount is not None and expense_amount != 0
            if has_income and has_expense:
                return None, "同一列同時有收入與支出金額"
            if not has_income and not has_expense:
                # Treat a non-empty row with no amount as an invalid accounting row.
                return None, "沒有收入或支出金額"
            if has_income:
                kind = "income"; amount = income_amount
                category = str(self._mapped_value(mapping, "income_category", values) or "").strip()
                summary = str(self._mapped_value(mapping, "income_summary", values) or "").strip()
            else:
                kind = "expense"; amount = expense_amount
                category = str(self._mapped_value(mapping, "expense_category", values) or "").strip()
                summary = str(self._mapped_value(mapping, "expense_summary", values) or "").strip()
            if not category:
                return None, "科目空白"
            if amount is None or amount < 1:
                return None, "金額錯誤"
            if amount > MAX_AMOUNT:
                return None, "金額最多7位數"

        if weighted_units(summary) > 40:
            return None, "摘要超過20個中文字或40個英數字元"
        return {
            "row_no": excel_row,
            "tx_date": tx_date,
            "account_name": account,
            "kind": kind,
            "category_name": category,
            "summary": summary,
            "amount": int(amount),
        }, None

    def _iter_rows(self, limit: int | None = None):
        ws = self._worksheet()
        if ws is None:
            return
        start = self._header_row_number() + 1
        count = 0
        for excel_row, values in enumerate(ws.iter_rows(min_row=start, values_only=True), start=start):
            record, error = self._normalized_row(excel_row, tuple(values))
            if record is None and error is None:
                continue
            yield excel_row, record, error
            count += 1
            if limit is not None and count >= limit:
                break

    def refresh_preview(self):
        self.preview.setRowCount(0)
        if self.workbook is None:
            return
        try:
            items = list(self._iter_rows(limit=20))
        except Exception as exc:
            self.preview.setRowCount(1)
            self.preview.setItem(0, 0, QTableWidgetItem(""))
            self.preview.setItem(0, 6, QTableWidgetItem(f"分析失敗：{exc}"))
            return
        for row_idx, (excel_row, record, error) in enumerate(items):
            self.preview.insertRow(row_idx)
            self.preview.setItem(row_idx, 0, QTableWidgetItem(str(excel_row)))
            if record:
                values = [
                    record["tx_date"], record["account_name"],
                    "收入" if record["kind"] == "income" else "支出",
                    record["category_name"], record["summary"], f"${record['amount']:,}",
                ]
                for col, value in enumerate(values, start=1):
                    self.preview.setItem(row_idx, col, QTableWidgetItem(str(value)))
                self.preview.item(row_idx, 6).setToolTip("可匯入")
            else:
                self.preview.setItem(row_idx, 6, QTableWidgetItem("錯誤：" + str(error)))

    @staticmethod
    def _combo_column(combo: QComboBox) -> int:
        data = combo.currentData()
        return int(data) if data is not None else -1

    def _required_mapping_ok(self) -> tuple[bool, str]:
        mode = self.mode_combo.currentData()
        mapping = self._current_mapping()
        fields = self.SPLIT_FIELDS if mode == self.MODE_SPLIT else self.UNIFIED_FIELDS
        missing = [label for key, label, required in fields if required and self._combo_column(mapping[key]) < 0]
        if mode == self.MODE_SPLIT:
            # At least one side must have category+amount mappings.
            inc_ok = self._combo_column(mapping["income_category"]) >= 0 and self._combo_column(mapping["income_amount"]) >= 0
            exp_ok = self._combo_column(mapping["expense_category"]) >= 0 and self._combo_column(mapping["expense_amount"]) >= 0
            if not inc_ok and not exp_ok:
                missing.append("至少一組收入或支出的科目＋金額")
        if missing:
            return False, "請先指定必要欄位：" + "、".join(missing)
        return True, ""

    def do_import(self):
        ok, message = self._required_mapping_ok()
        if not ok:
            QMessageBox.warning(self, APP_NAME, message)
            return
        records: list[dict] = []
        errors: list[str] = []
        try:
            for excel_row, record, error in self._iter_rows(limit=None):
                if record:
                    records.append(record)
                elif error:
                    errors.append(f"第 {excel_row} 列：{error}")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Excel 分析失敗：{exc}")
            return
        if not records:
            detail = "\n".join(errors[:12])
            QMessageBox.warning(self, APP_NAME, "沒有可匯入的有效資料。" + ("\n\n" + detail if detail else ""))
            return
        if errors:
            detail = "\n".join(errors[:10])
            if len(errors) > 10:
                detail += f"\n……另有 {len(errors)-10} 筆"
            answer = QMessageBox.question(
                self,
                APP_NAME,
                f"找到 {len(records)} 筆可匯入資料，另有 {len(errors)} 筆異常將略過。\n\n{detail}\n\n是否繼續？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel,
                QMessageBox.StandardButton.Cancel,
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
        else:
            if QMessageBox.question(
                self, APP_NAME, f"準備匯入 {len(records)} 筆資料，是否繼續？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel,
                QMessageBox.StandardButton.Cancel,
            ) != QMessageBox.StandardButton.Yes:
                return
        try:
            result = self.db.import_transactions(records, skip_duplicates=self.skip_duplicates.isChecked())
        except DatabaseError as exc:
            QMessageBox.warning(self, APP_NAME, str(exc))
            return
        self.imported_count = int(result.get("imported", 0))
        duplicates = int(result.get("duplicates", 0))
        db_errors = result.get("errors") or []
        text = f"匯入完成：{self.imported_count} 筆"
        if duplicates:
            text += f"\n重複略過：{duplicates} 筆"
        if errors or db_errors:
            text += f"\n異常略過：{len(errors) + len(db_errors)} 筆"
        QMessageBox.information(self, APP_NAME, text)
        self.accept()

    def done(self, result: int):  # noqa: N802
        try:
            if self.workbook is not None:
                self.workbook.close()
        except Exception:
            pass
        if self.temp_path is not None:
            try:
                self.temp_path.unlink(missing_ok=True)
            except Exception:
                pass
        super().done(result)
